"""Revisão estrutural de um .pptx gerado: o que dá para checar sem olhar. A revisão visual (olhar os PNGs) é do agente.

Erros: placeholder de texto vazio (aparece como "Clique para adicionar" ao editar).
Avisos: fonte abaixo de 12 pt, possível estouro de texto, imagem sem texto alternativo, mesmo layout três vezes seguidas.
"""
from __future__ import annotations

import math
from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.util import Emu

FONTE_MINIMA = 12
LIMITE_RODAPE_IN = 7.1  # texto abaixo disso é rodapé (nota de autoria, número): pode ser menor


def _formas(shapes):
    for sh in shapes:
        if sh.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from _formas(sh.shapes)
        else:
            yield sh


def _estouro(sh) -> str | None:
    """Estimativa grosseira: caracteres por linha pela largura da caixa e altura das linhas."""
    tf = sh.text_frame
    larg_pt = Emu(sh.width).pt if sh.width else 0
    alt_pt = Emu(sh.height).pt if sh.height else 0
    if larg_pt <= 0 or alt_pt <= 0:
        return None
    total = 0.0
    for p in tf.paragraphs:
        tam = next((r.font.size.pt for r in p.runs if r.font.size), None)
        if not tam or not p.text.strip():
            if not tam:
                return None  # tamanho vem do layout: não dá para estimar aqui
            continue
        cpl = max(1, int(larg_pt / (0.52 * tam)))
        total += math.ceil(len(p.text) / cpl) * tam * 1.2
    if total > alt_pt * 1.35 + 6:
        return f"texto pode não caber na caixa (precisa de ~{total / 72:.1f} pol, tem {alt_pt / 72:.1f})"
    return None


def revisar(pptx: Path) -> dict:
    prs = Presentation(str(pptx))
    erros: list = []
    avisos: list = []
    layouts = []
    for i, s in enumerate(prs.slides, 1):
        layouts.append(s.slide_layout.name)
        for ph in s.placeholders:
            if ph.has_text_frame and not ph.text_frame.text.strip():
                erros.append(f"slide {i}: placeholder de texto vazio ('{ph.name}'); remova ou preencha")
        for sh in _formas(s.shapes):
            if sh.shape_type == MSO_SHAPE_TYPE.PICTURE:
                if not (sh._element.xpath(".//p:cNvPr/@descr") or [""])[0]:
                    avisos.append(f"slide {i}: imagem '{sh.name}' sem texto alternativo")
            if sh.has_text_frame and sh.text_frame.text.strip():
                topo_in = Emu(sh.top).inches if sh.top is not None else 0
                for p in sh.text_frame.paragraphs:
                    for r in p.runs:
                        if r.font.size and r.text.strip() and r.font.size.pt < FONTE_MINIMA and topo_in < LIMITE_RODAPE_IN:
                            avisos.append(f"slide {i}: fonte de {r.font.size.pt:.0f} pt em '{r.text[:30]}' "
                                          f"(mínimo {FONTE_MINIMA}); encurte o texto em vez de reduzir a fonte")
                            break
                msg = _estouro(sh)
                if msg:
                    avisos.append(f"slide {i}: {msg} em '{sh.text_frame.text[:30]}'")
    for i in range(2, len(layouts)):
        if layouts[i] == layouts[i - 1] == layouts[i - 2]:
            avisos.append(f"slides {i - 1} a {i + 1}: o mesmo layout três vezes seguidas ('{layouts[i]}')")
    sem_notas = [i for i, s in enumerate(prs.slides, 1) if not (s.has_notes_slide and s.notes_slide.notes_text_frame.text.strip())]
    return {"slides": len(prs.slides), "erros": erros, "avisos": avisos, "slides_sem_notas": sem_notas}
