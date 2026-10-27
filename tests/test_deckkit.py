"""Testes do deckkit. Rodar: python -m unittest discover -s tests -v"""
import json
import sys
import tempfile
import unittest
from pathlib import Path

RAIZ = Path(__file__).resolve().parent.parent
SCRIPTS = RAIZ / "plugins" / "deck" / "scripts"
sys.path.insert(0, str(SCRIPTS))

from pptx import Presentation  # noqa: E402

from deckkit import marca as mod_marca  # noqa: E402
from deckkit import versoes  # noqa: E402
from deckkit.construtor import construir  # noqa: E402
from deckkit.deck import Deck  # noqa: E402
from deckkit.revisar import revisar  # noqa: E402
from deckkit.tipos import LAYOUTS, TIPOS, ErroDeSlide  # noqa: E402
from deckkit.validar import validar  # noqa: E402

EXEMPLOS = RAIZ / "plugins" / "deck" / "skills" / "deck" / "references" / "exemplos"
NEUTRO = RAIZ / "plugins" / "deck" / "marcas" / "neutro"


def _marca_temp(tmp: Path, **extra) -> mod_marca.Marca:
    pasta = tmp / "marca-teste"
    pasta.mkdir()
    dados = {"nome": "Teste", **extra}
    (pasta / "marca.json").write_text(json.dumps(dados), encoding="utf-8")
    return mod_marca.carregar(pasta)


class Modelo(unittest.TestCase):
    def test_construir_modelo_tem_os_nove_layouts(self):
        with tempfile.TemporaryDirectory() as t:
            m = _marca_temp(Path(t))
            arq = construir(m)
            prs = Presentation(str(arq))
            self.assertEqual([l.name for l in prs.slide_layouts], list(LAYOUTS.values()))
            self.assertTrue(arq.with_suffix(".potx").is_file())

    def test_marca_neutra_tem_modelo_pronto(self):
        self.assertTrue((NEUTRO / "modelo.pptx").is_file())


class Exemplos(unittest.TestCase):
    def test_exemplos_validam_geram_e_revisam_sem_erro(self):
        m = mod_marca.carregar(NEUTRO)
        for arq in sorted(EXEMPLOS.glob("*.json")):
            with self.subTest(exemplo=arq.stem):
                deck = json.loads(arq.read_text(encoding="utf-8"))
                r = validar(deck, m)
                self.assertEqual(r["erros"], [], r["erros"])
                d = Deck(m, projeto=deck.get("projeto", ""))
                for s in deck["slides"]:
                    d.adicionar(s)
                with tempfile.TemporaryDirectory() as t:
                    destino = Path(t) / "x.pptx"
                    d.salvar(destino)
                    rev = revisar(destino)
                    self.assertEqual(rev["erros"], [], rev["erros"])
                    self.assertEqual(rev["slides"], len(deck["slides"]))

    def test_todos_os_tipos_aparecem_nos_exemplos(self):
        usados = set()
        for arq in EXEMPLOS.glob("*.json"):
            usados |= {s["tipo"] for s in json.loads(arq.read_text(encoding="utf-8"))["slides"]}
        self.assertEqual(set(TIPOS) - usados, set(), "tipo sem exemplo")


class Validacao(unittest.TestCase):
    def setUp(self):
        self.m = mod_marca.carregar(NEUTRO)

    def _val(self, *slides, **marca_extra):
        return validar({"slides": list(slides)}, self.m)

    def test_texto_acima_do_limite_e_erro(self):
        r = self._val({"tipo": "decisao", "titulo": "x" * 80})
        self.assertTrue(any("limite" in e for e in r["erros"]), r)

    def test_tipo_desconhecido_e_campo_faltando(self):
        r = self._val({"tipo": "foguete"}, {"tipo": "numero"})
        self.assertEqual(len(r["erros"]), 2, r)
        self.assertIn("desconhecido", r["erros"][0])
        self.assertIn("faltam", r["erros"][1])

    def test_tres_do_mesmo_tipo_gera_aviso(self):
        s = {"tipo": "lista", "titulo": "T", "itens": ["a", "b"]}
        r = self._val(s, s, s)
        self.assertTrue(any("três do mesmo tipo" in a for a in r["avisos"]), r)

    def test_erro_de_um_slide_nao_esconde_os_outros(self):
        ok = {"tipo": "lista", "titulo": "ok", "itens": ["a"]}
        r = self._val({"tipo": "lista", "titulo": "y" * 90, "itens": ["a"]}, ok, {"tipo": "numero"})
        self.assertEqual(len(r["erros"]), 2, r)

    def test_travessao_so_e_erro_se_a_marca_proibe(self):
        s = {"tipo": "lista", "titulo": "A — B", "itens": ["x"]}
        self.assertEqual(validar({"slides": [s]}, self.m)["erros"], [])
        with tempfile.TemporaryDirectory() as t:
            m2 = _marca_temp(Path(t), estilo={"sem_travessao": True, "vedados": ["segredo"]})
            self.assertTrue(any("travessão" in e for e in validar({"slides": [s]}, m2, NEUTRO / "modelo.pptx")["erros"]))
            s2 = {"tipo": "lista", "titulo": "tem segredo aqui", "itens": ["x"]}
            self.assertTrue(any("vedado" in e for e in validar({"slides": [s2]}, m2, NEUTRO / "modelo.pptx")["erros"]))

    def test_numero_sem_fonte_so_avisa_se_a_marca_exige(self):
        s = {"tipo": "numero", "numero": "12.400", "rotulo": "solicitações em 2026"}
        self.assertFalse(any("fonte" in a for a in validar({"slides": [s]}, self.m)["avisos"]))
        with tempfile.TemporaryDirectory() as t:
            m2 = _marca_temp(Path(t), estilo={"exigir_fonte": True})
            r = validar({"slides": [s]}, m2, NEUTRO / "modelo.pptx")
            self.assertTrue(any("fonte" in a for a in r["avisos"]), r)

    def test_campo_fonte_vira_nota(self):
        d = Deck(self.m)
        d.adicionar({"tipo": "lista", "titulo": "T", "itens": ["a"], "fonte": "relatório de set"})
        textos = [sh.text_frame.text for sh in d.prs.slides[0].shapes if sh.has_text_frame]
        self.assertTrue(any("Fonte: relatório de set" in x for x in textos), textos)

    def test_notas_do_apresentador(self):
        d = Deck(self.m)
        d.adicionar({"tipo": "lista", "titulo": "T", "itens": ["a"], "notas": "texto do apresentador"})
        self.assertEqual(d.prs.slides[0].notes_slide.notes_text_frame.text, "texto do apresentador")

    def test_nenhum_placeholder_vazio_com_poucos_dados(self):
        d = Deck(self.m)
        d.adicionar({"tipo": "numero", "numero": "5", "rotulo": "x"})
        d.adicionar({"tipo": "decisao", "titulo": "T"})
        for s in d.prs.slides:
            for ph in s.placeholders:
                if ph.has_text_frame:
                    self.assertTrue(ph.text_frame.text.strip(), f"{s.slide_layout.name}: {ph.name}")


class Marcas(unittest.TestCase):
    def test_cores_derivadas_e_invalidas(self):
        m = mod_marca.carregar(NEUTRO)
        self.assertEqual(len(m.painel), 6)
        self.assertEqual(len(m.petalas), 4)
        with tempfile.TemporaryDirectory() as t:
            with self.assertRaises(mod_marca.ErroDeMarca):
                _marca_temp(Path(t), cores={"primaria": "azul"})

    def test_arquivo_de_ativo_inexistente_e_erro_claro(self):
        with tempfile.TemporaryDirectory() as t:
            m = _marca_temp(Path(t), logo="nao-existe.png")
            with self.assertRaises(mod_marca.ErroDeMarca):
                m.ativo("logo")

    def test_instalar_e_localizar(self):
        with tempfile.TemporaryDirectory() as t:
            antigo = mod_marca.CASA
            mod_marca.CASA = Path(t) / "casa"
            try:
                origem = Path(t) / "minha-marca"
                origem.mkdir()
                (origem / "marca.json").write_text('{"nome": "Minha"}', encoding="utf-8")
                destino = mod_marca.instalar(origem)
                self.assertTrue((destino / "marca.json").is_file())
                self.assertEqual(mod_marca.localizar("minha-marca").nome, "Minha")
                with self.assertRaises(mod_marca.ErroDeMarca):
                    mod_marca.localizar("nao-existe")
            finally:
                mod_marca.CASA = antigo


class Versoes(unittest.TestCase):
    def test_versao_nova_nunca_sobrescreve_e_detecta_edicao(self):
        with tempfile.TemporaryDirectory() as t:
            pasta = Path(t)
            d1 = versoes.proximo_destino(pasta, "[Proj] Deck")
            d1.write_bytes(b"um")
            versoes.registrar(d1, 3)
            d2 = versoes.proximo_destino(pasta, "[Proj] Deck")
            self.assertNotEqual(d1, d2)
            self.assertTrue(d2.name.endswith("v2.pptx"))
            self.assertEqual(versoes.verificar(pasta), [])
            d1.write_bytes(b"editado")
            av = versoes.verificar(pasta)
            self.assertEqual(len(av), 1)
            self.assertIn("EDITADO", av[0]["situacao"])

    def test_nome_sem_caracteres_proibidos_nem_espaco_duplo(self):
        self.assertEqual(versoes.limpar_nome("[Proj] Deck: v/1"), "Proj Deck v 1")


if __name__ == "__main__":
    unittest.main()
